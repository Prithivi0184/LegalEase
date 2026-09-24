from io import BytesIO
import html
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Mm

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle,
)
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)


# =========================================================
# TEXT SANITIZATION
# =========================================================

def sanitize_text(text: str) -> str:
    """
    Clean generated legal document text before formatting.
    """

    if not text:
        return ""

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2212": "-",
        "\u2026": "...",
        "\u00a0": " ",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    # Remove unsupported control characters
    text = re.sub(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]",
        "",
        text,
    )

    # Normalize excessive spaces
    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    # Normalize excessive blank lines
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


# =========================================================
# HTML PREVIEW FORMATTER
# =========================================================

def format_html_preview(text: str) -> str:
    """
    Convert legal document text into a styled HTML preview.
    """

    text = sanitize_text(text)

    if not text:
        return """
        <div class="legal-empty">
            No document content available.
        </div>
        """

    html_blocks = []

    for line in text.splitlines():

        line = line.strip()

        if not line:
            html_blocks.append(
                '<div class="legal-spacer"></div>'
            )
            continue

        escaped_line = html.escape(line)

        # Numbered legal headings:
        # 1. INTRODUCTION
        # 2. PARTIES
        # 10. TERMINATION
        numbered_heading = re.match(
            r"^\d+\.\s+.+",
            line,
        )

        # Common section headings
        upper_heading = (
            len(line) <= 100
            and line.upper() == line
            and any(char.isalpha() for char in line)
        )

        # Bullet points
        bullet_line = re.match(
            r"^[-•*]\s+(.+)",
            line,
        )

        if numbered_heading or upper_heading:

            html_blocks.append(
                f'<div class="legal-heading">'
                f'{escaped_line}'
                f'</div>'
            )

        elif bullet_line:

            bullet_text = html.escape(
                bullet_line.group(1)
            )

            html_blocks.append(
                f'<div class="legal-bullet">'
                f'• {bullet_text}'
                f'</div>'
            )

        else:

            html_blocks.append(
                f'<div class="legal-paragraph">'
                f'{escaped_line}'
                f'</div>'
            )

    content = "\n".join(html_blocks)

    return f"""
<style>

.legal-preview-card {{
    background-color: #151515;
    border: 1px solid #333333;
    border-radius: 12px;
    padding: 30px;
    height: 650px;
    overflow-y: auto;
    color: #eeeeee;
    font-family: Georgia, "Times New Roman", serif;
    font-size: 16px;
    line-height: 1.7;
    box-sizing: border-box;
}}

.legal-preview-title {{
    color: #ffffff;
    font-family: Arial, sans-serif;
    font-size: 22px;
    font-weight: 700;
    margin-bottom: 18px;
}}

.legal-heading {{
    color: #ffffff;
    font-family: Georgia, "Times New Roman", serif;
    font-size: 17px;
    font-weight: 700;
    margin-top: 20px;
    margin-bottom: 8px;
}}

.legal-paragraph {{
    margin-bottom: 10px;
}}

.legal-bullet {{
    margin-left: 20px;
    margin-bottom: 7px;
}}

.legal-spacer {{
    height: 8px;
}}

.legal-empty {{
    padding: 20px;
    color: #888888;
}}

</style>

<div class="legal-preview-title">
    ⚖️ LegalEase Document Preview
</div>

<div class="legal-preview-card">
    {content}
</div>
"""


# =========================================================
# TXT EXPORT
# =========================================================

def create_txt(document_text: str) -> bytes:

    cleaned_text = sanitize_text(
        document_text
    )

    return cleaned_text.encode(
        "utf-8"
    )


# =========================================================
# DOCX EXPORT
# =========================================================

def create_docx(document_text: str) -> bytes:

    document_text = sanitize_text(
        document_text
    )

    document = Document()

    section = document.sections[0]

    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    # -----------------------------------------------------
    # Title
    # -----------------------------------------------------

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        "LEGALEASE"
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)


    # -----------------------------------------------------
    # Subtitle
    # -----------------------------------------------------

    subtitle = document.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = subtitle.add_run(
        "AI-Powered Legal Document"
    )

    run.font.name = "Times New Roman"
    run.font.size = Pt(11)


    document.add_paragraph()


    # -----------------------------------------------------
    # Document Content
    # -----------------------------------------------------

    for line in document_text.splitlines():

        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        # Detect headings
        is_heading = (
            bool(
                re.match(
                    r"^\d+\.\s+.+",
                    line,
                )
            )
            or (
                len(line) <= 100
                and line.upper() == line
                and any(
                    char.isalpha()
                    for char in line
                )
            )
        )

        if is_heading:

            run = paragraph.add_run(
                line
            )

            run.bold = True
            run.font.name = (
                "Times New Roman"
            )
            run.font.size = Pt(12)

        else:

            run = paragraph.add_run(
                line
            )

            run.font.name = (
                "Times New Roman"
            )
            run.font.size = Pt(12)


    # -----------------------------------------------------
    # AI Notice
    # -----------------------------------------------------

    notice = document.add_paragraph()

    notice.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = notice.add_run(
        "AI-generated draft. Review by a qualified "
        "legal professional before signing or relying "
        "upon this document."
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


    output = BytesIO()

    document.save(output)

    return output.getvalue()


# =========================================================
# PDF EXPORT
# =========================================================

def create_pdf(document_text: str) -> bytes:

    document_text = sanitize_text(
        document_text
    )

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()


    # -----------------------------------------------------
    # Styles
    # -----------------------------------------------------

    title_style = ParagraphStyle(
        "LegalEaseTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        alignment=TA_CENTER,
        spaceAfter=8,
    )


    subtitle_style = ParagraphStyle(
        "LegalEaseSubtitle",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        alignment=TA_CENTER,
        spaceAfter=8,
    )


    heading_style = ParagraphStyle(
        "LegalEaseHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=15,
        spaceBefore=8,
        spaceAfter=6,
    )


    body_style = ParagraphStyle(
        "LegalEaseBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        spaceAfter=6,
    )


    bullet_style = ParagraphStyle(
        "LegalEaseBullet",
        parent=body_style,
        leftIndent=15,
        firstLineIndent=-8,
        spaceAfter=5,
    )


    notice_style = ParagraphStyle(
        "LegalEaseNotice",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=12,
        alignment=TA_CENTER,
        spaceBefore=12,
    )


    story = []


    # -----------------------------------------------------
    # PDF Header
    # -----------------------------------------------------

    story.append(
        Paragraph(
            "LEGALEASE",
            title_style,
        )
    )

    story.append(
        Paragraph(
            "AI-Powered Legal Document",
            subtitle_style,
        )
    )

    story.append(
        Spacer(
            1,
            8,
        )
    )


    # -----------------------------------------------------
    # PDF Content
    # -----------------------------------------------------

    for line in document_text.splitlines():

        line = line.strip()

        if not line:

            story.append(
                Spacer(
                    1,
                    5,
                )
            )

            continue


        safe_line = html.escape(
            line
        )


        # Heading
        is_heading = (
            bool(
                re.match(
                    r"^\d+\.\s+.+",
                    line,
                )
            )
            or (
                len(line) <= 100
                and line.upper() == line
                and any(
                    char.isalpha()
                    for char in line
                )
            )
        )


        # Bullet
        bullet_match = re.match(
            r"^[-•*]\s+(.+)",
            line,
        )


        if is_heading:

            story.append(
                Paragraph(
                    safe_line,
                    heading_style,
                )
            )

        elif bullet_match:

            bullet_text = html.escape(
                bullet_match.group(1)
            )

            story.append(
                Paragraph(
                    f"• {bullet_text}",
                    bullet_style,
                )
            )

        else:

            story.append(
                Paragraph(
                    safe_line,
                    body_style,
                )
            )


    # -----------------------------------------------------
    # AI Notice
    # -----------------------------------------------------

    story.append(
        Paragraph(
            "AI DRAFT NOTICE: This document was generated "
            "by LegalEase using generative AI. It should be "
            "reviewed by a qualified legal professional "
            "before being signed or relied upon.",
            notice_style,
        )
    )


    # -----------------------------------------------------
    # Build PDF
    # -----------------------------------------------------

    document.build(
        story,
        onFirstPage=_add_pdf_header_footer,
        onLaterPages=_add_pdf_header_footer,
    )


    return output.getvalue()


# =========================================================
# PDF HEADER / FOOTER
# =========================================================

def _add_pdf_header_footer(
    canvas,
    doc,
):

    canvas.saveState()

    width, height = A4


    # Header
    canvas.setFont(
        "Helvetica-Bold",
        9,
    )

    canvas.drawString(
        20 * mm,
        height - 15 * mm,
        "LegalEase",
    )


    # Footer
    canvas.setFont(
        "Helvetica",
        8,
    )

    canvas.drawRightString(
        width - 20 * mm,
        10 * mm,
        f"Page {doc.page}",
    )


    canvas.restoreState()


# =========================================================
# DOCUMENT FORMATTING ALIASES
# =========================================================

def format_docx(
    text: str,
    doc_type: str = "",
) -> bytes:

    return create_docx(text)


def format_pdf(
    text: str,
    doc_type: str = "",
) -> bytes:

    return create_pdf(text)